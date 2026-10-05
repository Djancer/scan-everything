import argparse
import logging
import tempfile
from pathlib import Path

from bot.application import build_application, build_services
from config import Settings, get_settings


def smoke_check() -> None:
    """Exercise extraction, mock metadata, JSON/TXT and SQLite without secrets."""
    from docx import Document

    with tempfile.TemporaryDirectory(prefix="document_bot_smoke_") as temp_dir:
        root = Path(temp_dir)
        source = root / "sample_invoice.docx"
        document = Document()
        document.add_heading("Sample Invoice", level=1)
        document.add_paragraph("Invoice amount: 19.99 EUR")
        document.add_paragraph("Offline smoke check")
        document.save(source)

        settings = Settings(
            llm_provider="mock",
            ocr_provider="disabled",
            data_dir=root / "data",
            database_path=root / "data" / "documents.db",
        )
        pipeline, database = build_services(settings)
        result = pipeline.process(source, source.name, telegram_user_id=1)
        assert Path(result.document.text_path).exists()
        assert Path(result.document.metadata_path).exists()
        assert database.search(1, "invoice")
        print(f"SMOKE CHECK OK: {result.document.id} ({result.document.metadata.document_type})")


def main() -> None:
    parser = argparse.ArgumentParser(description="Telegram AI Document Bot")
    parser.add_argument("--check", action="store_true", help="run offline smoke check")
    args = parser.parse_args()
    if args.check:
        smoke_check()
        return

    settings = get_settings()
    logging.basicConfig(
        level=getattr(logging, settings.log_level.upper(), logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s: %(message)s",
    )
    application = build_application(settings)
    application.run_polling(allowed_updates=["message"])


if __name__ == "__main__":
    main()

