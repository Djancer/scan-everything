from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import pymupdf as fitz
from docx import Document
from PIL import Image

from .ocr import OCRProvider, OCRUnavailableError


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".jpg", ".jpeg", ".png"}


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(slots=True)
class ExtractionResult:
    text: str
    used_ocr: bool = False
    warnings: list[str] = field(default_factory=list)


class DocumentTextExtractor:
    def __init__(self, ocr: OCRProvider) -> None:
        self.ocr = ocr

    def extract(self, path: Path) -> ExtractionResult:
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(f"Unsupported file type: {suffix or 'unknown'}")
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        return self._extract_image(path)

    def _extract_docx(self, path: Path) -> ExtractionResult:
        document = Document(path)
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append("\t".join(values))
        return ExtractionResult(text="\n".join(blocks))

    def _extract_image(self, path: Path) -> ExtractionResult:
        try:
            return ExtractionResult(text=self.ocr.extract_path(path), used_ocr=True)
        except OCRUnavailableError as exc:
            return ExtractionResult(text="", used_ocr=False, warnings=[str(exc)])

    def _extract_pdf(self, path: Path) -> ExtractionResult:
        pages: list[str] = []
        warnings: list[str] = []
        used_ocr = False
        with fitz.open(path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                if text:
                    pages.append(text)
                    continue
                # A page without a usable text layer is rasterized and OCR'd once.
                try:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                        text = self.ocr.extract_image(image)
                    used_ocr = True
                    pages.append(text)
                except OCRUnavailableError as exc:
                    warnings.append(f"Page {page_number}: {exc}")
                    pages.append("")
        return ExtractionResult(text="\n\n".join(pages).strip(), used_ocr=used_ocr, warnings=warnings)
